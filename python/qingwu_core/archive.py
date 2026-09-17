from __future__ import annotations

import html
import json
import shutil
import tempfile
import zipfile
from pathlib import Path
from typing import Any

from .errors import ConflictError, ValidationError
from .util import json_dumps, safe_relative_path, sha256_file, unique_path, utc_now


def _safe_name(value: str, fallback: str = "未命名事务") -> str:
    result = "".join("_" if char in '<>:"/\\|?*' else char for char in value).strip(" .")
    return result or fallback


def build_manifest(bundle: dict[str, Any], warnings: list[dict[str, Any]]) -> dict[str, Any]:
    facts = {}
    for key, field in bundle["facts"].items():
        if field.get("sensitive") or field.get("status") != "confirmed":
            continue
        facts[key] = {
            "label": field.get("label", key),
            "value": field.get("value"),
            "status": "confirmed",
        }
    return {
        "schema_version": 1,
        "app": {"name": "Qingwu", "version": "0.1.0"},
        "affair": {
            "id": bundle["id"], "title": bundle["title"],
            "template": {"id": bundle["template_id"], "version": bundle["template_version"]},
            "fact_version": bundle["current_fact_version"], "ai_used": bool(bundle["ai_used"]),
        },
        "confirmed_facts": facts,
        "completed_tasks": [
            {"title": task["title"], "stage": task["stage"], "completed": bool(task["completed"])}
            for task in bundle["tasks"] if task["completed"]
        ],
        "final_drafts": [
            {"document_id": draft["document_id"], "title": draft["title"], "version": draft["version"],
             "status": draft["status"], "fact_version": draft["fact_version"]}
            for draft in bundle["drafts"]
        ],
        "materials": [
            {"slot_id": item.get("slot_id"), "group_instance_id": item.get("group_instance_id"),
             "name": item["original_name"], "size_bytes": item["size_bytes"], "sha256": item["sha256"]}
            for item in bundle["materials"]
        ],
        "unresolved_warnings": warnings,
        "exported_at": utc_now(),
    }


def checklist_html(manifest: dict[str, Any]) -> str:
    affair = manifest["affair"]
    facts = "".join(
        f"<tr><td>{html.escape(item['label'])}</td><td>{html.escape(str(item.get('value') or '未确认'))}</td>"
        f"<td>{html.escape(item['status'])}</td></tr>"
        for item in manifest["confirmed_facts"].values()
    )
    tasks = "".join(f"<li>{html.escape(item['title'])}</li>" for item in manifest["completed_tasks"]) or "<li>暂无</li>"
    materials = "".join(
        f"<tr><td>{html.escape(item['name'])}</td><td>{html.escape(item.get('slot_id') or '待归类')}</td>"
        f"<td><code>{item['sha256']}</code></td></tr>" for item in manifest["materials"]
    ) or "<tr><td colspan='3'>暂无材料</td></tr>"
    warnings = "".join(
        f"<li><strong>{html.escape(item.get('code', 'warning'))}</strong>：{html.escape(item['message'])}</li>"
        for item in manifest["unresolved_warnings"]
    ) or "<li>无</li>"
    return f"""<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8"><title>{html.escape(affair['title'])} · 归档清单</title>
<style>body{{font-family:"Microsoft YaHei",sans-serif;max-width:980px;margin:40px auto;color:#27312d;padding:0 24px}}h1{{color:#185a43}}.meta{{color:#66736d}}table{{border-collapse:collapse;width:100%;margin:12px 0 28px}}th,td{{border:1px solid #dbe2de;padding:9px;text-align:left}}th{{background:#edf5f1}}code{{font-size:11px;word-break:break-all}}@media print{{body{{margin:0}}}}</style></head>
<body><h1>{html.escape(affair['title'])}</h1><p class="meta">轻务归档清单 · 模板 {html.escape(affair['template']['id'])} {html.escape(affair['template']['version'])} · 事实版本 v{affair['fact_version']}</p>
<h2>事实底稿</h2><table><tr><th>字段</th><th>值</th><th>状态</th></tr>{facts}</table>
<h2>已完成待办</h2><ul>{tasks}</ul><h2>材料</h2><table><tr><th>文件</th><th>槽位</th><th>SHA-256</th></tr>{materials}</table>
<h2>未解决警告</h2><ul>{warnings}</ul><p class="meta">导出时间：{html.escape(manifest['exported_at'])} · 使用过 AI：{'是' if affair['ai_used'] else '否'}</p></body></html>"""


def _write_pdf(path: Path, manifest: dict[str, Any]) -> bool:
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.pdfgen import canvas
    except ImportError:
        return False
    font_name = "Helvetica"
    for candidate in (Path("C:/Windows/Fonts/msyh.ttc"), Path("C:/Windows/Fonts/simhei.ttf")):
        if candidate.exists():
            try:
                pdfmetrics.registerFont(TTFont("QingwuChinese", str(candidate), subfontIndex=0))
                font_name = "QingwuChinese"
                break
            except Exception:
                continue
    page_width, page_height = A4
    doc = canvas.Canvas(str(path), pagesize=A4)
    y = page_height - 50

    def line(text: str, size: int = 10, gap: int = 18) -> None:
        nonlocal y
        if y < 55:
            doc.showPage()
            y = page_height - 50
        doc.setFont(font_name, size)
        doc.drawString(44, y, text[:100])
        y -= gap

    line(manifest["affair"]["title"], 18, 28)
    line(f"轻务归档清单 · 事实版本 v{manifest['affair']['fact_version']}", 10, 24)
    line("事实底稿", 14, 24)
    for item in manifest["confirmed_facts"].values():
        line(f"{item['label']}：{item.get('value') or '未确认'}（{item['status']}）")
    line("材料", 14, 24)
    for item in manifest["materials"]:
        line(f"{item['name']} · {item.get('slot_id') or '待归类'}")
        line(f"SHA-256: {item['sha256']}", 7, 13)
    line("未解决警告", 14, 24)
    for warning in manifest["unresolved_warnings"]:
        line(f"[{warning.get('code', 'warning')}] {warning['message']}")
    doc.save()
    return True


def _write_docx(path: Path, title: str, body: str) -> bool:
    try:
        from docx import Document
    except ImportError:
        return False
    document = Document()
    document.add_heading(title, level=1)
    for paragraph in body.split("\n"):
        document.add_paragraph(paragraph)
    document.save(path)
    return True


def export_affair(bundle: dict[str, Any], warnings: list[dict[str, Any]], output: str | Path,
                  *, allow_warnings: bool = True) -> dict[str, Any]:
    blocking = [item for item in warnings if item.get("severity") == "error"]
    if blocking and not allow_warnings:
        raise ConflictError("事务仍有阻断问题，不能导出")
    output_path = Path(output).expanduser().resolve()
    if output_path.suffix.lower() != ".zip":
        output_path.mkdir(parents=True, exist_ok=True)
        output_path = output_path / f"{_safe_name(bundle['title'])}.zip"
    else:
        output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path = unique_path(output_path)
    manifest = build_manifest(bundle, warnings)
    template = bundle["template"]
    material_slots = {slot["id"]: slot for slot in template.get("material_slots", [])}
    with tempfile.TemporaryDirectory(prefix="qingwu-export-") as temporary:
        root = Path(temporary) / _safe_name(bundle["title"])
        (root / "documents").mkdir(parents=True)
        (root / "materials").mkdir(parents=True)
        drafts_dir = root / "drafts"
        drafts_dir.mkdir(parents=True)
        copied = []
        for material in bundle["materials"]:
            source = Path(material["source_path"])
            if not source.is_file():
                continue
            if sha256_file(source) != material["sha256"]:
                raise ConflictError(f"导出时材料内容发生变化：{material['original_name']}")
            slot = material_slots.get(material.get("slot_id"))
            relative_dir = safe_relative_path(slot["output_path"]) if slot else Path("materials/待归类")
            if material.get("group_instance_id"):
                relative_dir = Path("materials") / _safe_name(material["group_instance_id"]) / relative_dir.name
            target_dir = root / relative_dir
            target_dir.mkdir(parents=True, exist_ok=True)
            target = unique_path(target_dir / _safe_name(material["original_name"], "material"))
            shutil.copy2(source, target)
            if sha256_file(target) != material["sha256"]:
                raise ConflictError(f"复制后材料哈希不一致：{material['original_name']}")
            copied.append(str(target.relative_to(root)).replace("\\", "/"))
        docx_available = True
        for draft in bundle["drafts"]:
            stem = _safe_name(f"{draft['title']}-v{draft['version']}")
            extension = ".md" if draft["kind"].startswith("article") else ".txt"
            (drafts_dir / f"{stem}{extension}").write_text(draft["rendered_body"], encoding="utf-8")
            if draft["kind"].startswith("article"):
                article_html = "<html lang='zh-CN'><meta charset='utf-8'><body><pre>" + html.escape(draft["rendered_body"]) + "</pre></body></html>"
                (drafts_dir / f"{stem}.html").write_text(article_html, encoding="utf-8")
                docx_available = _write_docx(drafts_dir / f"{stem}.docx", draft["title"], draft["rendered_body"]) and docx_available
        if not docx_available:
            warnings.append({"code": "docx_unavailable", "severity": "warning", "message": "运行环境缺少 python-docx，未生成 DOCX 文案。"})
            manifest = build_manifest(bundle, warnings)
        (root / "qingwu-manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
        html_text = checklist_html(manifest)
        (root / "qingwu-checklist.html").write_text(html_text, encoding="utf-8")
        pdf_created = _write_pdf(root / "qingwu-checklist.pdf", manifest)
        if not pdf_created:
            warnings.append({"code": "pdf_unavailable", "severity": "warning", "message": "运行环境缺少 reportlab，未生成 PDF 清单。"})
            manifest = build_manifest(bundle, warnings)
            (root / "qingwu-manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
            (root / "qingwu-checklist.html").write_text(checklist_html(manifest), encoding="utf-8")
        with zipfile.ZipFile(output_path, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=6) as archive:
            for path in sorted(root.rglob("*")):
                if path.is_file():
                    archive.write(path, path.relative_to(root).as_posix())
    return {
        "path": str(output_path), "sha256": sha256_file(output_path), "size_bytes": output_path.stat().st_size,
        "material_files": copied, "warning_count": len(warnings), "manifest": manifest,
    }
